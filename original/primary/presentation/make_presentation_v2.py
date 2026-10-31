from pathlib import Path
import csv
��om pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor

ROOT = Path(__file__).resolve().parents[1]
FIG = ROOT / "results" / "figures"
CSV = ROOT / "results" / "scenario_summary_v2.csv"
OUT = ROOT / "presentation" / "DataCenter_Heat_Reuse_v0_2_RU.pptx"

with CSV.open(encoding="utf-8-sig") as f:
    rows = list(csv.DictReader(f))

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

NAVY = RGBColor(19,38,58)
BLUE = RGBColor(35,104,158)
CYAN = RGBColor(40,145,170)
LIGHT = RGBColor(244,247,250)
DARK = RGBColor(31,41,55)
MUTED = RGBColor(90,105,120)
WHITE = RGBColor(255,255,255)
GREEN = RGBColor(39,128,95)

def title(s,t,sub=None):
    b=s.shapes.add_textbox(Inches(.65),Inches(.32),Inches(12),Inches(.62))
    p=b.text_frame.paragraphs[0]; p.text=t; p.font.size=Pt(26); p.font.bold=True; p.font.color.rgb=NAVY
    if sub:
        b=s.shapes.add_textbox(Inches(.67),Inches(.98),Inches(11.8),Inches(.38))
        p=b.text_frame.paragraphs[0]; p.text=sub; p.font.size=Pt(11.5); p.font.color.rgb=MUTED

def footer(s,n):
    ln=s.shapes.add_shape(MSO_SHAPE.RECTANGLE,Inches(.65),Inches(7.16),Inches(12),Inches(.012))
    ln.fill.solid(); ln.fill.fore_color.rgb=RGBColor(210,218,226); ln.line.fill.background()
    b=s.shapes.add_textbox(Inches(11.9),Inches(7.19),Inches(.55),Inches(.18))
    p=b.text_frame.paragraphs[0]; p.text=str(n); p.font.size=Pt(8); p.font.color.rgb=MUTED; p.alignment=PP_ALIGN.RIGHT

def bullets(s,items,x=.8,y=1.55,w=11.6,h=5.2,size=18):
    b=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h)); tf=b.text_frame; tf.word_wrap=True
    for i,it in enumerate(items):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph()
        p.text=it; p.font.size=Pt(size); p.font.color.rgb=DARK; p.space_after=Pt(10)
    return b

def image(s,name,x,y,w):
    s.shapes.add_picture(str(FIG"name),Inches(x),Inches(y),width=Inches(w))

def card(s,x,y,w,label,val,note=""):
    sh=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(1.15))
    sh.fill.solid(); sh.fill.fore_color.rgb=LIGHT; sh.line.color.rgb=RGBColor(205,215,225)
    b=s.shapes.add_textbox(Inches(x+.18),Inches(y+.13),Inches(w-.36),Inches(.85)); tf=b.text_frame
    p=tf.paragraphs[0]; p.text=val; p.font.size=Pt(23); p.font.bold=True; p.font.color.rgb=BLUE
    p=tf.add_paragraph(); p.text=label; p.font.size=Pt(10.5); p.font.color.rgb=DARK
    if note:
        p=tf.add_paragraph(); p.text=note; p.font.size=Pt(7.5); p.font.color.rgb=MUTED

# 1
s=prs.slides.add_slide(prs.slide_layouts[6])
s.background.fill.solid(); s.background.fill.fore_color.rgb=NAVY
b=s.shapes.add_textbox(Inches(.9),Inches(1.2),Inches(11.6),Inches(2.3)); tf=b.text_frame
p=tf.paragraphs[0]; p.text="Утилизация тепла дата-центра"; p.font.size=Pt(39); p.font.bold=True; p.font.color.rgb=WHITE
p=tf.add_paragraph(); p.text="Модульная модель дC → теплообменник → тепловой насос → накопитель → теплосеть"; p.font.size=Pt(20); p.font.color.rgb=RGBColor(191,219,254); p.space_before=Pt(10)
p=tf.add_paragraph(); p.text="MATLAB/Simulink research platform · v0.2 · 8760 h"; p.font.size=Pt(14); p.font.color.rgb=RGBColor(203,213,225); p.space_before=Pt(16)

# 2