from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / "presentation" / "DataCenter_Heat_Reuse_v0_2_RU.pptx"
OUT = ROOT / "presentation" / "DataCenter_Heat_Reuse_v0_3_RU.pptx"
FIG = ROOT / "results" / "figures" / "14_uncertainty_v3.png"

prs = Presentation(SRC)
for slide in prs.slides:
    for shape in slide.shapes:
        if not hasattr(shape, "text_frame"):
            continue
        for p in shape.text_frame.paragraphs:
            for run in p.runs:
                if "v0.2" in run.text:
                    run.text = run.text.replace("v0.2", "v0.3")

NAVY = RGBColor(19, 38, 58)
BLUE = RGBColor(35, 104, 158)
DARK = RGBColor(31, 41, 55)
MUTED = RGBColor(90, 105, 120)
LIGHT = RGBColor(244, 247, 250)

def add_title(slide, text, subtitle=None):
    box = slide.shapes.add_textbox(Inches(0.65), Inches(0.28), Inches(12.0), Inches(0.62))
    p = box.text_frame.paragraphs[0]
    p.text = text
    p.font.size = Pt(25)
    p.font.bold = True
    p.font.color.rgb = NAVY
    if subtitle:
        box = slide.shapes.add_textbox(Inches(0.67), Inches(0.92), Inches(11.7), Inches(0.35))
        p = box.text_frame.paragraphs[0]
        p.text = subtitle
        p.font.size = Pt(11)
        p.font.color.rgb = MUTED

def add_bullets(slide, items, x, y, w, h, size=17):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = item
        p.font.size = Pt(size)
        p.font.color.rgb = DARK
        p.space_after = Pt(9)
    return box

def footer(slide, number):
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.65), Inches(7.12), Inches(12), Inches(0.01))
    line.fill.solid(); line.fill.fore_color.rgb = RGBColor(215, 222, 229); line.line.fill.background()
    box = slide.shapes.add_textbox(Inches(11.9), Inches(7.15), Inches(0.5), Inches(0.2))
    p = box.text_frame.paragraphs[0]; p.text = str(number); p.font.size = Pt(8); p.font.color.rgb = MUTED; p.alignment = PP_ALIGN.RIGHT

# Uncertainty slide
s = prs.slides.add_slide(prs.slide_layouts[6])
add_title(s, "11. Устойчивость результата к допущениям", "200 сценарных прогонов вокруг Warm45 + 100 MWh storage")
s.shapes.add_picture(str(FIG), Inches(0.75), Inches(1.35), width=Inches(7.35))
add_bullets(s, [
    "Median COP: 6.25; P05-P95: 5.45-6.67.",
    "Median heat coverage: 81.9%; P05-P95: 77.8-85.4%.",
    "Median OPEX saving: 2.47 MEUR/year при сценарных ценах.",
    "Median simple payback: 5.23 years; P05-P95: 3.75-9.80 years.",
    "Интервалы показывают sensitivity reference-case, а не статистическую гарантию для реального объекта."
], 8.35, 1.55, 4.25, 4.9, 15)
footer(s, len(prs.slides))

# Roadmap/evidence slide
s = prs.slides.add_slide(prs.slide_layouts[6])
add_title(s, "12. Что уже доказано моделью и что нужно дальше", "Граница между reference study и calibrated engineering case")
add_bullets(s, [
    "Проверено: энергетический баланс, SOC/COP limits и направленные физические зависимости.",
    "Показано: temperature lift - центральная связь между cooling architecture, DH temperature и electricity use.",
    "Показано: storage улучшает temporal matching, но его размер должен оптимизироваться по CAPEX и сезонности.",
    "Следующий шаг: реальные weather/load/supply-return/DC telemetry + manufacturer heat-pump map.",
    "После калибровки: LP/MILP/MPC dispatch, seasonal storage, hydraulics, NPV/ownership и Bratislava case."
], 0.85, 1.55, 7.2, 4.9, 16)
# evidence card
card = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(8.0), Inches(1.55), Inches(4.45), Inches(4.7))
card.fill.solid(); card.fill.fore_color.rgb = LIGHT; card.line.color.rgb = RGBColor(205,215,225)
box = s.shapes.add_textbox(Inches(8.25), Inches(1.8), Inches(3.9), Inches(4.15))
tf = box.text_frame
for i, txt in enumerate([
    "Литературные опорные точки",
    "Wahlroos et al. (2017): Espoo, annual DH dispatch, 20-60 MW waste heat.",
    "Jurgens et al. (2024): Frankfurt, 144 GWh/a, 37 MWth HP + 20 MWth gas, 97.5% coverage case.",
    "Hiltunen & Syri (2021): Espoo, роль low-temperature waste heat и HP.",
    "Verzino et al. (2026): Varna/THUNDER, exergy + thermochemical storage + 3GDH."
]):
    p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
    p.text = txt; p.font.color.rgb = NAVY if i == 0 else DARK
    p.font.bold = (i == 0); p.font.size = Pt(15 if i == 0 else 9.6); p.space_after = Pt(8)
footer(s, len(prs.slides))

prs.save(OUT)
print(OUT)
