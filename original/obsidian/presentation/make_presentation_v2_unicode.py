from pathlib import Path
import csv
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor
R=Path(__file__).resolve().parents[1]; F=R/'results'/'figures'; C=R/'results'/'scenario_summary_v2.csv'; O=R/'presentation'/'DataCenter_Heat_Reuse_v0_2_RU.pptx'
with C.open(encoding='utf-8-sig') as f: rows=list(csv.DictReader(f))
prs=Presentation(); prs.slide_width=Inches(13.333); prs.slide_height=Inches(7.5)
NAVY=RGBColor(19,38,58); BLUE=RGBColor(35,104,158); LIGHT=RGBColor(244,247,250); DARK=RGBColor(31,41,55); MUTED=RGBColor(90,105,120); WHITE=RGBColor(255,255,255)
def T(s,t,sub=''):
 b=s.shapes.add_textbox(Inches(.65),Inches(.3),Inches(12),Inches(.65)); p=b.text_frame.paragraphs[0]; p.text=t; p.font.size=Pt(27); p.font.bold=True; p.font.color.rgb=NAVY
 if sub:
  b=s.shapes.add_textbox(Inches(.67),Inches(.98),Inches(11.8),Inches(.4)); p=b.text_frame.paragraphs[0]; p.text=sub; p.font.size=Pt(11); p.font.color.rgb=MUTED
def B(s,items,x=.85,y=1.55,w=11.5,h=5,size=18):
 b=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h)); tf=b.text_frame; tf.word_wrap=True
 for i,a in enumerate(items):
  p=tf.paragraphs[0] if i==0 else tf.add_paragraph(); p.text=a; p.font.size=Pt(size); p.font.color.rgb=DARK; p.space_after=Pt(12)
def I(s,n,x,y,w): s.shapes.add_picture(str(F/n),Inches(x),Inches(y),width=Inches(w))
def foot(s,n):
 b=s.shapes.add_textbox(Inches(11.9),Inches(7.15),Inches(.55),Inches(.2)); p=b.text_frame.paragraphs[0]; p.text=str(n); p.font.size=Pt(8); p.font.color.rgb=MUTED; p.alignment=PP_ALIGN.RIGHT
# 1
s=prs.slides.add_slide(prs.slide_layouts[6]); s.background.fill.solid(); s.background.fill.fore_color.rgb=NAVY
b=s.shapes.add_textbox(Inches(.9),Inches(1.25),Inches(11.6),Inches(2.2)); tf=b.text_frame
p=tf.paragraphs[0]; p.text="\u0423\u0442\u0438\u043b\u0438\u0437\u0430\u0446\u0438\u044f \u0442\u0435\u043f\u043b\u0430 \u0434\u0430\u0442\u0430-\u0446\u0435\u043d\u0442\u0440\u0430"; p.font.size=Pt(39); p.font.bold=True; p.font.color.rgb=WHITE
p=tf.add_paragraph(); p.text="\u041c\u043e\u0434\u0443\u043b\u044c\u043d\u0430\u044f MATLAB/Simulink-\u043c\u043e\u0434\u0435\u043b\u044c \u0438 \u0433\u043e\u0434\u043e\u0432\u043e\u0439 \u044d\u043a\u0441\u043f\u0435\u0440\u0438\u043c\u0435\u043d\u0442"; p.font.size=Pt(20); p.font.color.rgb=RGBColor(191,219,254); p.space_before=Pt(10)
p=tf.add_paragraph(); p.text='research platform v0.2 | 8760 h'; p.font.size=Pt(14); p.font.color.rgb=RGBColor(203,213,225); p.space_before=Pt(18)
#2
s=prs.slides.add_slide(prs.slide_layouts[6]); T(s,"1. \u0418\u0441\u0441\u043b\u0435\u0434\u043e\u0432\u0430\u0442\u0435\u043b\u044c\u0441\u043a\u0438\u0439 \u0432\u043e\u043f\u0440\u043e\u0441")
B(s,["\u041a\u0430\u043a \u0442\u0435\u043c\u043f\u0435\u0440\u0430\u0442\u0443\u0440\u0430 \u043e\u0442\u0445\u043e\u0434\u044f\u0449\u0435\u0433\u043e \u0442\u0435\u043f\u043b\u0430, \u0442\u0435\u043c\u043f\u0435\u0440\u0430\u0442\u0443\u0440\u0430 \u0442\u0435\u043f\u043b\u043e\u0441\u0435\u0442\u0438 \u0438 \u043d\u0430\u043a\u043e\u043f\u0438\u0442\u0435\u043b\u044c \u0432\u043b\u0438\u044f\u044e\u0442 \u043d\u0430 COP, \u043f\u043e\u043a\u0440\u044b\u0442\u0438\u0435 \u0441\u043f\u0440\u043e\u0441\u0430, \u0433\u0430\u0437\u043e\u0432\u044b\u0439 \u0440\u0435\u0437\u0435\u0440\u0432, \u0441\u0442\u043e\u0438\u043c\u043e\u0441\u0442\u044c \u0438 CO2?","\u0426\u0435\u043b\u044c: \u0440\u0430\u0441\u0448\u0438\u0440\u044f\u0435\u043c\u0430\u044f \u043f\u043b\u0430\u0442\u0444\u043e\u0440\u043c\u0430, \u043a\u043e\u0442\u043e\u0440\u0443\u044e \u043c\u043e\u0436\u043d\u043e \u043a\u0430\u043b\u0438\u0431\u0440\u043e\u0432\u0430\u0442\u044c \u043f\u043e\u0434 Bratislava / Frankfurt / Espoo."],.9,1.7,11.2,4.4,20); foot(s,2)
#3
s=prs.slides.add_slide(prs.slide_layouts[6]); T(s,"2. \u041b\u0438\u0442\u0435\u0440\u0430\u0442\u0443\u0440\u043d\u0430\u044f \u043e\u0441\u043d\u043e\u0432\u0430")
B(s,["Espoo: \u0433\u043e\u0434\u043e\u0432\u044b\u0435 \u0441\u0446\u0435\u043d\u0430\u0440\u0438\u0438 \u0443\u0442\u0438\u043b\u0438\u0437\u0430\u0446\u0438\u0438 \u0442\u0435\u043f\u043b\u0430.","Frankfurt: 144 GWh/a, 37 MWth HP + 20 MWth gas + storage.","Hiltunen & Syri: \u0446\u0435\u043d\u043d\u043e\u0441\u0442\u044c liquid cooling \u0438 \u0431\u043e\u043b\u0435\u0435 \u0432\u044b\u0441\u043e\u043a\u043e\u0439 \u0442\u0435\u043c\u043f\u0435\u0440\u0430\u0442\u0443\u0440\u044b \u0438\u0441\u0442\u043e\u0447\u043d\u0438\u043a\u0430.","Varna / THUNDER: storage \u0438 exergy \u043a\u0430\u043a \u0441\u043b\u0435\u0434\u0443\u044e\u0449\u0438\u0439 \u0443\u0440\u043e\u0432\u0435\u043d\u044c \u043c\u043e\u0434\u0435\u043b\u0438."],.9,1.6,11.1,4.8,18); foot(s,3)
#4 architecture
s=prs.slides.add_slide(prs.slide_layouts[6]); T(s,"3. \u041c\u043e\u0434\u0443\u043b\u044c\u043d\u0430\u044f \u0430\u0440\u0445\u0438\u0442\u0435\u043a\u0442\u0443\u0440\u0430 v0.2")
labels=['IT load','DC source','Heat exchanger','Heat pump','Dispatch','District heating','Backup']; xs=[.45,2.25,4.05,5.9,7.7,9.5,11.3]
for lab,x in zip(labels,xs):
 sh=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,Inches(x),Inches(2.25),Inches(1.45),Inches(.9)); sh.fill.solid(); sh.fill.fore_color.rgb=LIGHT; sh.line.color.rgb=BLUE
 b=s.shapes.add_textbox(Inches(x+.06),Inches(2.5),Inches(1.3),Inches(.35)); p=b.text_frame.paragraphs[0]; p.text=lab; p.font.size=Pt(12); p.font.bold=True; p.font.color.rgb=DARK; p.alignment=PP_ALIGN.CENTER
B(s,["\u041a\u0430\u0436\u0434\u044b\u0439 \u0444\u0438\u0437\u0438\u0447\u0435\u0441\u043a\u0438\u0439 \u0431\u043b\u043e\u043a \u0438\u043c\u0435\u0435\u0442 \u043e\u0442\u0434\u0435\u043b\u044c\u043d\u044b\u0439 \u0438\u043d\u0442\u0435\u0440\u0444\u0435\u0439\u0441 \u0438 \u043c\u043e\u0436\u0435\u0442 \u0431\u044b\u0442\u044c \u0437\u0430\u043c\u0435\u043d\u0451\u043d \u0431\u0435\u0437 \u043f\u0435\u0440\u0435\u043f\u0438\u0441\u044b\u0432\u0430\u043d\u0438\u044f \u0432\u0441\u0435\u0439 \u043c\u043e\u0434\u0435\u043b\u0438."],.9,4.5,11.3,1.1,16); foot(s,4)
#5 math + plot
s=prs.slides.add_slide(prs.slide_layouts[6]); T(s,"4. \u041c\u0430\u0442\u0435\u043c\u0430\u0442\u0438\u0447\u0435\u0441\u043a\u043e\u0435 \u044f\u0434\u0440\u043e")
B(s,["\u0418\u0441\u0442\u043e\u0447\u043d\u0438\u043a: Qsrc = \u03b7rec*PIT. \u0422\u0435\u043f\u043b\u043e\u043e\u0431\u043c\u0435\u043d\u043d\u0438\u043a: QHX = min(\u03b5HX*Qsrc,Qmax).","\u0422\u0435\u043f\u043b\u043e\u0432\u043e\u0439 \u043d\u0430\u0441\u043e\u0441: COP ~ \u03b7C*Th/(Th-Tc).","\u041d\u0430\u043a\u043e\u043f\u0438\u0442\u0435\u043b\u044c: E(k+1)=E(k)(1-\u03bb\u0394t)+\u03b7cQc\u0394t-Qd\u0394t/\u03b7d.","\u042d\u043a\u0441\u0435\u0440\u0433\u0438\u044f: ExQ=Q(1-T0/T)."],.8,1.55,6.0,4.8,16); I(s,'09_lift_cop_coverage_v2.png',7.0,1.55,5.6); foot(s,5)
#6 yearly
s=prs.slides.add_slide(prs.slide_layouts[6]); T(s,"5. \u0413\u043e\u0434\u043e\u0432\u043e\u0439 \u044d\u043a\u0441\u043f\u0435\u0440\u0438\u043c\u0435\u043d\u0442 v0.2","10 MW IT | 8760 h | synthetic Bratislava-like weather")
I(s,'08_scenario_comparison_v2.png',2.3,1.35,8.7)
foot(s,6)
#7 sensitivity
s=prs.slides.add_slide(prs.slide_layouts[6]); T(s,"6. \u0427\u0443\u0432\u0441\u0442\u0432\u0438\u0442\u0435\u043b\u044c\u043d\u043e\u0441\u0442\u044c Tsource x Tsupply")
I(s,'11_cop_map_v2.png',.65,1.45,5.9); I(s,'12_coverage_map_v2.png',6.75,1.45,5.9)
B(s,["\u041b\u0435\u0432\u0430\u044f \u043a\u0430\u0440\u0442\u0430 \u043f\u043e\u043a\u0430\u0437\u044b\u0432\u0430\u0435\u0442 COP, \u043f\u0440\u0430\u0432\u0430\u044f - \u043f\u043e\u043a\u0440\u044b\u0442\u0438\u0435 \u0441\u043f\u0440\u043e\u0441\u0430. \u041e\u043f\u0442\u0438\u043c\u0443\u043c\u044b \u043d\u0435 \u043e\u0431\u044f\u0437\u0430\u043d\u044b \u0441\u043e\u0432\u043f\u0430\u0434\u0430\u0442\u044c."],.9,6.0,11.4,.7,13); foot(s,7)
#8 storage
s=prs.slides.add_slide(prs.slide_layouts[6]); T(s,"7. \u041d\u0430\u043a\u043e\u043f\u0438\u0442\u0435\u043b\u044c: \u043f\u043e\u043b\u0435\u0437\u043d\u043e\u0441\u0442\u044c \u0438 diminishing returns")
I(s,'13_storage_tradeoff_v2.png',1.1,1.4,8.2)
B(s,["\u041d\u0435\u0431\u043e\u043b\u044c\u0448\u043e\u0439 storage \u0441\u0433\u043b\u0430\u0436\u0438\u0432\u0430\u0435\u0442 \u0447\u0430\u0441\u044b \u0438 \u0441\u0443\u0442\u043a\u0438.","\u0414\u043b\u044f \u0441\u0435\u0437\u043e\u043d\u043d\u043e\u0433\u043e \u043f\u0435\u0440\u0435\u043d\u043e\u0441\u0430 \u043d\u0443\u0436\u0435\u043d \u0434\u0440\u0443\u0433\u043e\u0439 \u043c\u0430\u0441\u0448\u0442\u0430\u0431 \u0438 \u043e\u0442\u0434\u0435\u043b\u044c\u043d\u0430\u044f \u043c\u043e\u0434\u0435\u043b\u044c \u043f\u043e\u0442\u0435\u0440\u044c."],9.55,1.8,2.8,3.8,14); foot(s,8)
#9 validation
s=prs.slides.add_slide(prs.slide_layouts[6]); T(s,"8. \u0412\u0430\u043b\u0438\u0434\u0430\u0446\u0438\u044f")
B(s,["\u0410\u0432\u0442\u043e\u0442\u0435\u0441\u0442\u044b \u043f\u0440\u043e\u0432\u0435\u0440\u044f\u044e\u0442 \u0431\u0430\u043b\u0430\u043d\u0441 \u0442\u0435\u043f\u043b\u043e\u0441\u0435\u0442\u0438, \u0433\u0440\u0430\u043d\u0438\u0446\u044b SOC, \u0434\u0438\u0430\u043f\u0430\u0437\u043e\u043d COP, \u043d\u0435\u043e\u0442\u0440\u0438\u0446\u0430\u0442\u0435\u043b\u044c\u043d\u044b\u0435 \u043f\u043e\u0442\u043e\u043a\u0438 \u0438 \u043e\u0433\u0440\u0430\u043d\u0438\u0447\u0435\u043d\u0438\u0435 \u0434\u043e\u0441\u0442\u0443\u043f\u043d\u043e\u0433\u043e \u0442\u0435\u043f\u043b\u0430.","v0.1 \u0441\u043e\u0445\u0440\u0430\u043d\u0451\u043d \u043a\u0430\u043a regression baseline.","v1.0 \u043f\u043e\u0442\u0440\u0435\u0431\u0443\u0435\u0442 \u0440\u0435\u0430\u043b\u044c\u043d\u044b\u0445 weather/load/temperature/telemetry \u0434\u0430\u043d\u043d\u044b\u0445."],.9,1.7,11.2,4.6,19); foot(s,9)
#10 roadmap
s=prs.slides.add_slide(prs.slide_layouts[6]); T(s,"9. Roadmap")
B(s,["v0.3: \u0442\u0435\u043f\u043b\u043e\u043e\u0431\u043c\u0435\u043d\u043d\u0438\u043a UA/LMTD \u0438\u043b\u0438 \u03b5-NTU, \u0433\u0438\u0434\u0440\u0430\u0432\u043b\u0438\u043a\u0430, manufacturer HP map.","v0.4: CAPEX/NPV/LCOH, Monte-Carlo, Morris/Sobol.","v0.5: LP/MILP/MPC dispatch \u0441 \u043f\u0440\u043e\u0433\u043d\u043e\u0437\u043e\u043c \u043f\u043e\u0433\u043e\u0434\u044b, \u0446\u0435\u043d\u044b \u0438 \u043d\u0430\u0433\u0440\u0443\u0437\u043a\u0438.","v1.0: calibrated case + telemetry + Simscape + Telegram/app."],.9,1.7,11.2,4.7,19); foot(s,10)
prs.save(O); print(O)
